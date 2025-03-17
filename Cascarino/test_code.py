import openpyxl
import pprint

def get_section_data(order_of_data):
    # Location of modified section data file
    # Note that the naming of the profile needs to match that in RFEM
    path = "C:\\Users\\casca\\Desktop\Development\\RFEM_Python_Client\\Reference\Steel Profiles\\steel-profiles-south-africa.xlsx"

    wb_obj = openpyxl.load_workbook(path, data_only=True)
    sheet_obj = wb_obj.active

    # IPE and UBs
    # Position of data for IPEs and UBs in excel sheet
    start_row = 11
    end_row = 63

    I_section_data = []
    for i in range(start_row, end_row+1):
        cell_obj = sheet_obj.cell(row=i, column=5)
        if cell_obj.value != None:

            I_section_data.append({
                "name" : cell_obj.value,
                "mass" : float(sheet_obj.cell(row=i, column=6).value)
            })

    # UCs
    # Position of data for UCs in excel sheet
    start_row = 80
    end_row = 101

    H_section_data = []
    for i in range(start_row, end_row+1):
        cell_obj = sheet_obj.cell(row=i, column=5)
        if cell_obj.value != None:

            H_section_data.append({
                "name" : cell_obj.value,
                "mass" : float(sheet_obj.cell(row=i, column=6).value)
            })
    section_data = []
    for item in order_of_data:
        match item:
            case 'I-Sections':
                section_data.append(I_section_data)
            case 'H-Sections':
                section_data.append(H_section_data)

    return section_data

section_data = get_section_data(['H-Sections', 'I-Sections'])

list1 = sorted(section_data[0], key=lambda d: d['mass'])
pprint.pprint(list1)