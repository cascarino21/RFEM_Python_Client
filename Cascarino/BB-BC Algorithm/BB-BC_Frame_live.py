import random as rd

import time
import pprint
import openpyxl

from numpy import random
import matplotlib.pyplot as plt

import RFEM
from RFEM.enums import ObjectTypes
from RFEM.initModel import Model
from RFEM.BasicObjects.section import Section
from RFEM.BasicObjects.member import Member
from RFEM.Results.designOverview import *
from RFEM.Results.resultTables import ResultTables, ConvertResultsToListOfDct
from RFEM.Tools.GetObjectNumbersByType import GetObjectNumbersByType
from RFEM.Tools.centreOfGravityAndObjectInfo import ObjectsInfo

def get_section_data(order_of_data):
    # Location of modified section data file
    # Note that the naming of the profile needs to match that in RFEM
    path = "C:\\Users\\casca\\Desktop\Development\\RFEM_Python_Client\\Reference\Steel Profiles\\steel-profiles-south-africa.xlsx"

    wb_obj = openpyxl.load_workbook(path, data_only=True)
    sheet_obj = wb_obj.active

    # IPE and UBs
    # Position of data for IPEs and UBs in excel sheet
    start_row = 11
    end_row = 63

    I_section_data = []
    for i in range(start_row, end_row+1):
        cell_obj = sheet_obj.cell(row=i, column=5)
        if cell_obj.value != None:

            I_section_data.append({
                "name" : cell_obj.value,
                "mass" : float(sheet_obj.cell(row=i, column=6).value)
            })
    I_section_data = sorted(I_section_data, key= lambda d: d['mass'])

    # UCs
    # Position of data for UCs in excel sheet
    start_row = 80
    end_row = 101

    H_section_data = []
    for i in range(start_row, end_row+1):
        cell_obj = sheet_obj.cell(row=i, column=5)
        if cell_obj.value != None:

            H_section_data.append({
                "name" : cell_obj.value,
                "mass" : float(sheet_obj.cell(row=i, column=6).value)
            })
    H_section_data = sorted(H_section_data, key= lambda d: d['mass'])

    # Equal Angles
    # Position of data for EAs in excel sheet
    start_row = 138
    end_row = 195

    L_section_data = []
    for i in range(start_row, end_row+1):
        cell_obj = sheet_obj.cell(row=i, column=5)
        if cell_obj.value != None:

            L_section_data.append({
                "name" : cell_obj.value,
                "mass" : float(sheet_obj.cell(row=i, column=6).value)
            })
    L_section_data = sorted(L_section_data, key= lambda d: d['mass'])

    # CHSs
    # Position of data for CHSs in excel sheet
    start_row = 244
    end_row = 367

    CHS_section_data = []
    for i in range(start_row, end_row+1):
        cell_obj = sheet_obj.cell(row=i, column=5)
        if cell_obj.value != None:

            CHS_section_data.append({
                "name" : cell_obj.value,
                "mass" : round(float(sheet_obj.cell(row=i, column=6).value), 2)
            })
    CHS_section_data = sorted(CHS_section_data, key= lambda d: d['mass'])

    section_data = []
    for item in order_of_data:
        match item:
            case 'I-Sections':
                section_data.append(I_section_data)
            case 'H-Sections':
                section_data.append(H_section_data)
            case 'CHS':
                section_data.append(CHS_section_data)
            case 'Equal Angles':
                section_data.append(L_section_data)
            case 'Not-optimisable':
                section_data.append([0])

    return section_data

def calculate_weights(genome: list, section_lib: list[list], length_per_sections: list[float], output_required: bool = True):
    weight  = 0
    for gene, length, sections in zip(genome, length_per_sections, section_lib):
            if len(sections) == 1:
                break
            else:
                section_mass = sections[gene]["mass"]
                weight = weight + length * section_mass

    if output_required:
        print(f"Genome: {genome} Total Weight = {round(weight,2)} ")
    return round(weight,2)

def generate_genome(section_data: {list[list]}):
    return [rd.randint(0,len(section_list)-1) for section_list in section_data]

def initialize_population(population_size: int, section_data: list[list]):
    return [generate_genome(section_data) for _ in range(population_size)]

# -- Functions related to RFEM --
def connect_to_RFEM(file_name):
    model = RFEM.initModel.Model(False, model_name=file_name)
    print("Connected!")

def get_results() -> list:
    RFEM.initModel.Calculate_all()
    return ResultTables.SteelDesignDesignRatiosMembersBySection(True)

def run_structural_analysis(genome: list[int], section_lib: list[list]):
    print(f"Check genome: {genome}", end="...   ")
    begin = time.time()

    # Update the sections
    for gene, i, sections in zip(genome, range(len(genome)), section_lib):
        if len(sections) == 1:
            break
        else:
            section_name = sections[gene]["name"]
            Section(no= i+1, name=section_name, material_no=1)

    RFEM.initModel.Calculate_all()
    elapsed = time.time() - begin
    print(f"...calculation took {round(elapsed,1)} seconds")
    return ResultTables.SteelDesignDesignRatiosMembersBySection(True)

def get_RFEM_geometry():
    # Check how this works for a section list with a blank in between
    # Might need to modify this using max(list_of_sections)
    member_list = GetObjectNumbersByType(ObjectTypes.E_OBJECT_TYPE_MEMBER)
    sections_in_model = GetObjectNumbersByType(ObjectTypes.E_OBJECT_TYPE_SECTION)
    section_names = []
    # Set total length for each section to 0
    lengths = [0.0 for _ in range(len(sections_in_model))]

    for member in member_list:
        member_dict = Model.clientModel.service.get_member(member)
        lengths[member_dict["section_start"]-1] = lengths[member_dict["section_start"]-1] + member_dict["length"]

    for section in sections_in_model:
        section_dict = Model.clientModel.service.get_section(section)
        section_names.append(section_dict["name"])

    return sections_in_model, lengths, section_names

def change_sections(sectionList: list, section_number: int):
    for section in sectionList:
        Section(no=section_number, name=section)

# End of functions related to RFEM --------------------------

def get_classification_list(member_list: list):
    classification_list = []
    for member in member_list:
        section_name = member.split(' | ')[0]
        components = section_name.split(' ')
        if len(components) > 2:
            classification = components[0]+' '+components[1]
        else:
            classification = components[0]
        match classification:
            case 'UB':
                classification_list.append('I-Sections')
            case 'IPE':
                classification_list.append('I-Sections')
            case 'UC':
                classification_list.append('H-Sections')
            case 'CHS':
                classification_list.append('CHS')
            case 'L':
                classification_list.append('Equal Angles')
            case other:
                classification_list.append('Not-optimisable')

    return classification_list

def find_max_design_ratios(data: list) -> list:
    max_design_ratios = []

    for data_set, i in zip(data, range(len(data))):
        # print("Check: {i}")
        # print(data_set)
        max_design_ratios.append([])
        section_count = -1
        if data_set == "UBS Applied":
            max_design_ratios[i] = ([["UBS Applied", 1.0, 1.0]])
        else:
            for design_check in data_set:
                description = design_check['description']
                if isinstance(description, str):
                    section_count = section_count + 1
                    section_name = description.split(" | ")[0]
                    max_design_ratios[i].append([section_name, float(section_count)+1, 0.0 ])
                else:
                    design_ratio = design_check['design_ratio']
                    if isinstance(design_ratio, str):
                        if design_ratio == "Non-designable":
                            print(f"Note that {section_name} is Class 4 and Non-designable for a specific case")
                    else:
                        max_design_ratios[i][section_count][2] = max(max_design_ratios[i][section_count][2], round(design_ratio,2))
    return max_design_ratios

def penalised_objective(weight, max_design_ratios) -> float:
    multiplier = 1
    feasible = True
    for design_ratio in max_design_ratios:
        if design_ratio[2] >1:
            multiplier = (multiplier + design_ratio[2])
            feasible = False
    if feasible:
        penalised_weight = weight
    else:
        penalised_weight = weight*multiplier
    return penalised_weight

def generate_new_point(cog, generation, section_data, alpha = 1.0):
    new_point = []
    for cog_i,list in zip(cog,section_data):
        temp = int(round(cog_i + alpha*random.normal(0, 1)*len(list)/generation, 0))
        new_point.append(min(max(temp, 0), len(list)-1))
    return new_point

# Main Algorithm

connect_to_RFEM(file_name = "Model - Portal Frame")
sections_in_model, length_per_sections, section_names = get_RFEM_geometry()
classification_list = get_classification_list(member_list=section_names)
section_data = get_section_data(classification_list)

# Parameters
population_size = 5
generations = 5
all_calculated = []
all_data = []
saved_iterations = 0

# Create population
population = initialize_population(population_size=population_size, section_data=section_data)

gens_without_improvement = 0
temp_optimimum = []
current_best_weight = float("inf")

# plt.axis([0,generations, 0, 500000])

# Algorithm
for generation in range(1, generations+1):

    generation_data = []
    print(f"\nBeginning Calculation Generation: {generation}")

    weights = [calculate_weights(genome, section_data, length_per_sections, output_required=False) for genome in population]
    for genome, weight in zip(population, weights):
        # Check if the current arrangement has already been calculated to save computational time
        if genome in all_calculated:
            print(f"Genome {genome} already calculated")
            pos = all_calculated.index(genome)
            generation_data.append(all_data[pos])
            saved_iterations = saved_iterations + 1
        elif weight > current_best_weight:
            saved_iterations = saved_iterations + 1
            generation_data.append("UBS Applied")
        else:
            data = run_structural_analysis(genome = genome, section_lib=section_data)
            generation_data.append(data)
            all_data.append(data)
            all_calculated.append(genome)

    max_design_ratios = []
    max_design_ratios = find_max_design_ratios(data=generation_data)

    # Penalizing objectives
    penalised_weights = [penalised_objective(weight, max_design_ratio,) for weight, max_design_ratio in zip(weights, max_design_ratios)]

    sortedPopulation = [genome for _,genome in sorted(zip(penalised_weights,population), reverse=False)]
    sortedPenalizedWeights = sorted((penalised_weights), reverse=False)
    sortedWeights = [weight for _,weight in sorted(zip(penalised_weights,weights), reverse=False)]

    # Determine center of gravity
    cog = sortedPopulation[0]

    # # Debugging code ----------------------
    # print("\nResults:")
    # for genome, max_design_ratio, weight, penalised_weight in zip(population, max_design_ratios, weights, penalised_weights):
    #     print(f"Genome: {genome}: || Design Ratios: {max_design_ratio}, || Weight: {round(weight,2)} || Penalized Weight: {round(penalised_weight,2)} ")
    # # --------------------------------------

    # Checking if there is a change in the best solution from previous
    if sortedPopulation[0] == temp_optimimum:
        gens_without_improvement = gens_without_improvement + 1
    else:
        gens_without_improvement = 0
        temp_optimimum = sortedPopulation[0]

    # [0:1] only takes first value but it saves it as an array
    new_population = sortedPopulation[0:1]
    for _ in range(population_size-1):
        while True:
            temp_point = generate_new_point(cog=sortedPopulation[0], generation=generation, section_data=section_data)
            temp_weight = calculate_weights(temp_point, section_data, length_per_sections, output_required=False)
            if temp_weight < sortedPenalizedWeights[0]:
                break
        new_population.append(temp_point)
    population = new_population
    section_1, section_2 = [],[]
    for point in population:
        section_1.append(point[0])
        section_2.append(point[1])

    # Print best fitness in this generation
    current_best_weight = min(penalised_weights)

    # plt.scatter(generation, current_best_weight)
    # plt.pause(0.05)

    print(f"Generation {generation}, Solution: {population[0]}, Best penalized weight: {current_best_weight}, Actual Weight: {round(sortedWeights[0],2)}, Gens without improvement: {gens_without_improvement} ")

    if gens_without_improvement == 5:
        break


# plt.show()
print(f"Number of analyses saved = {saved_iterations}")
