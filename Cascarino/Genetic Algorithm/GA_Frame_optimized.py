import random
from typing import List

import time
import pprint
import openpyxl

import RFEM

from RFEM.enums import ObjectTypes
from RFEM.initModel import Model
from RFEM.BasicObjects.section import Section
from RFEM.BasicObjects.member import Member
from RFEM.Results.designOverview import *
from RFEM.Results.resultTables import ResultTables, ConvertResultsToListOfDct
from RFEM.Tools.GetObjectNumbersByType import GetObjectNumbersByType
from RFEM.Tools.centreOfGravityAndObjectInfo import ObjectsInfo

# Parameters
population_size = 2
generations = 1
mutation_rate = 0.3
mutation_count = 0
applied_moment = 140


def get_section_data(order_of_data):
    # Location of modified section data file
    # Note that the naming of the profile needs to match that in RFEM
    path = "C:\\Users\\casca\\Desktop\Development\\RFEM_Python_Client\\Reference\Steel Profiles\\steel-profiles-south-africa.xlsx"

    wb_obj = openpyxl.load_workbook(path, data_only=True)
    sheet_obj = wb_obj.active

    # IPE and UBs
    # Position of data for IPEs and UBs in excel sheet
    start_row = 11
    end_row = 63

    I_section_data = []
    for i in range(start_row, end_row+1):
        cell_obj = sheet_obj.cell(row=i, column=5)
        if cell_obj.value != None:

            I_section_data.append({
                "name" : cell_obj.value,
                "mass" : float(sheet_obj.cell(row=i, column=6).value)
            })

    # UCs
    # Position of data for UCs in excel sheet
    start_row = 80
    end_row = 101

    H_section_data = []
    for i in range(start_row, end_row+1):
        cell_obj = sheet_obj.cell(row=i, column=5)
        if cell_obj.value != None:

            H_section_data.append({
                "name" : cell_obj.value,
                "mass" : float(sheet_obj.cell(row=i, column=6).value)
            })
    section_data = []
    for item in order_of_data:
        match item:
            case 'I-Sections':
                section_data.append(I_section_data)
            case 'H-Sections':
                section_data.append(H_section_data)

    return section_data

def calculate_weights(genome: List, section_lib: List[List], length_per_sections: List[float], output_required: bool = True):
    weight  = 0
    for gene, length, sections in zip(genome, length_per_sections, section_lib):
            section_mass = sections[gene]["mass"]

            # Here I could add a check to see if this value is already greater than the total weight of the current optimum solution and I could potentially remove this item from the potential sections

            weight = weight + length * section_mass

    if output_required:
        print(f"Genome: {genome} Total Weight = {round(weight,2)} ")
    return round(weight,2)

def generate_genome(section_data: {List[List]}):
    return [random.randint(0,len(section_list)-1) for section_list in section_data]

def initialize_population(population_size: int, section_data: List[List]):
    return [generate_genome(section_data) for _ in range(population_size)]

# -- Functions related to RFEM --
def connect_to_RFEM():
    model = RFEM.initModel.Model(False, model_name="Model - Portal Frame.rf6")
    print("Connected!")

def get_results() -> list:
    RFEM.initModel.Calculate_all()
    # Get design ratios
    return ResultTables.SteelDesignDesignRatiosMembersBySection(True)

def run_structural_analysis(genome: list[int], section_lib: list[list]):
    print(f"Check genome: {genome}")
    begin = time.time()

    # Update the sections
    for gene, i, sections in zip(genome, range(len(genome)), section_lib):
        section_name = sections[gene]["name"]
        Section(no= i+1, name=section_name, material_no=1)

    RFEM.initModel.Calculate_all()

    elapsed = time.time() - begin
    print(f"Calculation took {round(elapsed,1)} seconds")
    # count = count +1

    return ResultTables.SteelDesignDesignRatiosMembersBySection(True)

def get_RFEM_geometry():
    # Check how this works for a section list with a blank in between
    # Might need to modify this using max(list_of_sections)

    list_of_members = GetObjectNumbersByType(ObjectTypes.E_OBJECT_TYPE_MEMBER)
    sections_in_model = GetObjectNumbersByType(ObjectTypes.E_OBJECT_TYPE_SECTION)

    # Set total length for each section to 0
    lengths = [0.0 for _ in range(len(sections_in_model))]

    for member in list_of_members:
        member_dict = Model.clientModel.service.get_member(member)
        lengths[member_dict["section_start"]-1] = lengths[member_dict["section_start"]-1] + member_dict["length"]

    return sections_in_model, lengths

def change_sections(sectionList: list, section_number: int):
    for section in sectionList:
        Section(no=section_number, name=section)

# End of functions connecting to RFEM --------------------------

def find_max_design_ratios(data: list) -> list:
    max_design_ratios = []
    for data_set, i in zip(data, range(len(data))):
        max_design_ratios.append([])
        section_count = -1
        for design_check in data_set:
            description = design_check['description']
            if isinstance(description, str):
                section_count = section_count + 1
                section_name = description.split(" | ")[0]
                max_design_ratios[i].append([section_name, float(section_count)+1, 0.0 ])
            else:
                design_ratio = design_check['design_ratio']
                if isinstance(design_ratio, str):
                    if design_ratio == "Non-designable":
                        print(f"Note that {section_name} is Class 4 and Non-designable for a specific case")
                else:
                    max_design_ratios[i][section_count][2] = max(max_design_ratios[i][section_count][2], round(design_ratio,2))
    return max_design_ratios

def get_min_weights(weights, max_design_ratios) -> float:
    min_weight_all = min(weights)
    min_weight_feas = max(weights)

    for weight, max_design_ratio in zip(weights, max_design_ratios):
        max_design_ratio_solution = 0
        for design_ratio in max_design_ratio:
            max_design_ratio_solution = max(max_design_ratio_solution, design_ratio[2])

        if ((max_design_ratio_solution <= 1.00) & (weight < min_weight_feas)):
            min_weight_feas = weight

    return min_weight_all, min_weight_feas

def penalised_objective(weight, max_design_ratios, min_weight_all, min_weight_feas) -> float:
    multiplier = 1
    feasible = True
    for design_ratio in max_design_ratios:
        if design_ratio[2] >1:
            multiplier = multiplier * design_ratio[2]
            feasible = False
    if feasible:
        penalised_weight = weight
    else:
        penalised_weight = weight + (min_weight_feas-min_weight_all)*(multiplier)
    return penalised_weight


def fitness_function(min_weight, max_weight, penalised_weight) -> float:
    fitness = min_weight + max_weight - penalised_weight
    return fitness

# Selection
def tournament_selection(population, fitnesses, k=3):
    selected_indices = random.sample(range(len(population)), k)
    return max(selected_indices, key=lambda i: fitnesses[i])

# Crossover
def single_point_crossover(parent1, parent2):
    point = random.randint(1, len(parent1) - 1)
    child1 = parent1[:point] + parent2[point:]
    child2 = parent2[:point] + parent1[point:]
    return child1, child2

# Mutation
def random_resetting(genome):
    chance = random.random()
    if(mutation_rate > chance):
        genome = generate_genome(section_data=section_data)
        # print("Mutation occured")

    return genome

# Genetic Algorithm
def genetic_algorithm(population, generations, mutation_rate):
    # results = []
    for _ in generations:
        for genome in population:
            this_result = run_structural_analysis(genome=genome)
            # results.append(this_result)
            print(this_result[0])

# Main Algorithm

# Need to later get user input to confirm the order of the section data in the list
section_data = get_section_data(['H-Sections', 'I-Sections'])

connect_to_RFEM()
sections_in_model, length_per_sections = get_RFEM_geometry()
# print(sections_in_model, length_per_sections)

# Create population
population_size = 10
population = initialize_population(population_size=population_size, section_data=section_data)

# genetic_algorithm(population=population, generations=1, mutation_rate=0.3)
generations = 10
all_calculated = []
all_data = []
# applied_moment_input = int(input("Applied moment (in kNm):"))

for generation in range(generations):

    generation_data = []
    print(f"Beginning Calculation Generation: {generation+1}")
    for genome in population:
        if genome in all_calculated:
            print(f"Genome {genome} already calculated")
            pos = all_calculated.index(genome)
            generation_data.append(all_data[pos])
        else:
            data = run_structural_analysis(genome = genome, section_lib=section_data)
            generation_data.append(data)
            all_data.append(data)
            all_calculated.append(genome)

    max_design_ratios = []
    max_design_ratios = find_max_design_ratios(data=generation_data)

    weights = []
    weights = [calculate_weights(genome, section_data, length_per_sections, output_required=False) for genome in population]

    # Penalizing objectives
    min_weight_all, min_weight_feas = get_min_weights(weights=weights, max_design_ratios=max_design_ratios)
    penalised_weights = [penalised_objective(weight, max_design_ratio, min_weight_all=min_weight_all, min_weight_feas=min_weight_feas) for weight, max_design_ratio in zip(weights, max_design_ratios)]
    # print(weights)

    # Calculating weights of all
    max_weight = max(penalised_weights)
    min_weight = min(penalised_weights)
    fitnesses = [fitness_function(min_weight=min_weight, max_weight=max_weight, penalised_weight=penalised_weight) for penalised_weight in penalised_weights]

    sortedPopulation = [genome for _,genome in sorted(zip(fitnesses,population), reverse=True)]
    # sortedUtilization = [utilization for _,utilization in sorted(zip(fitnesses,utilizations), reverse=True)]
    sortedPenalizedWeights = [weight for _,weight in sorted(zip(fitnesses,penalised_weights), reverse=True)]
    sortedWeights = [weight for _,weight in sorted(zip(fitnesses,weights), reverse=True)]

    # Debugging code ----------------------
    print("\nResults:")
    for genome, max_design_ratio, weight, penalised_weight, fit in zip(population, max_design_ratios, weights, penalised_weights, fitnesses):
        print(f"Genome: {genome}: || Design Ratios: {max_design_ratio}, || Weight: {round(weight,2)} || Penalized Weight: {round(penalised_weight,2)} || Fitness: {round(fit,2)}")

    # --------------------------------------

    new_population = sortedPopulation[0:1]
    for genome in sortedPopulation:
        if genome != new_population[0]:
            new_population.append(genome)
            break
    if len(new_population) ==1:
        print(f"All items in the popoulation are the same")
        break

    for _ in range((population_size // 2) -1):  # Two offspring per iteration
        parent1 = population[tournament_selection(population, fitnesses)]
        # Ensure that both parents are not the same
        while True:
            parent2 = population[tournament_selection(population, fitnesses)]
            if parent1 != parent2:
                break
        child1, child2 = single_point_crossover(parent1, parent2)
        child1 = random_resetting(child1)
        child2 = random_resetting(child2)
        new_population.extend([child1, child2])

    # Replace old population with new
    population = new_population

    # Print best fitness in this generation
    best_fitness = max(fitnesses)
    print(f"Generation {generation}, Best genome: {sortedPopulation[0]} has weight of {round(sortedPenalizedWeights[0],2)} and fitness of {round(best_fitness,2)}")


